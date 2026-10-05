"""HydroScope Water Research 1.0 — an evidence-first environmental research workbench.

Run: python -m streamlit run app.py
Python 3.12. No supplied artwork is bundled. No key is needed
for water-data adapters; an optional single LLM call interprets saved results.
See README.md and METHODS.md for coverage, attribution and scientific limits.
"""
from __future__ import annotations

import base64
import calendar
import hashlib
import html
import io
import json
import math
import os
import re
import threading
import time
import zipfile
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from urllib.parse import urlparse

import folium
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
import requests
import streamlit as st
from pyproj import Geod, Transformer
from requests.adapters import HTTPAdapter
from shapely.geometry import Point, Polygon, mapping, shape
from shapely.geometry.polygon import orient
from shapely.ops import transform as shape_transform, unary_union
from urllib3.util.retry import Retry

ROOT = Path(__file__).resolve().parent
VERSION = "hydroscope-1.0.0"
GEOD = Geod(ellps="WGS84")
TEAL, VIOLET, GOLD = "#5FE1C3", "#A78BFA", "#F4C97A"
COLORS = [TEAL, VIOLET, GOLD, "#76B8FF", "#F18BB8"]
MAX_SAT_KM2, MAX_SAT_PIXELS = 250.0, 600_000
SAT_COLLECTION = "sentinel-2-c1-l2a"
MODULES = ["Satellite", "River outlook", "Marine outlook"]
PAGES = ["Overview", "Water study", "Satellite water maps", "Water research", "River & marine", "AI interpretation", "Reports & sources"]
FIELD_COLUMNS = ["site", "date", "latitude", "longitude", "chlorophyll_ug_l", "secchi_m", "total_phosphorus_ug_l", "dissolved_oxygen_mg_l", "ph", "temperature_c", "turbidity_ntu", "notes"]
SOURCES = {
    "Satellite": "https://github.com/Element84/earth-search",
    "River outlook": "https://open-meteo.com/en/docs/flood-api",
    "Marine outlook": "https://open-meteo.com/en/docs/marine-weather-api",
    "Field observations": "https://www.nalms.org/secchidipin/monitoring-methods/trophic-state-equations/",
}


class DataError(Exception):
    """A user-readable validation or provider failure."""


def utc_now():
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def secret(name, default=""):
    try:
        return str(st.secrets.get(name, os.getenv(name, default)))
    except (FileNotFoundError, st.errors.StreamlitSecretNotFoundError):
        return os.getenv(name, default)


def request_json(url, params=None, payload=None):
    """Two attempts at most; API credentials never appear in errors/logs."""
    retry = Retry(total=0 if urlparse(url).hostname == "nominatim.openstreetmap.org" else 1, connect=1, read=0, backoff_factor=0.4,
                  status_forcelist=[429, 500, 502, 503, 504],
                  allowed_methods=frozenset(["GET"]), respect_retry_after_header=False)
    with requests.Session() as session:
        session.mount("https://", HTTPAdapter(max_retries=retry))
        headers = {"User-Agent": "HydroScope/1.0 water research dashboard",
                   "Accept": "application/json"}
        try:
            response = session.request("POST" if payload is not None else "GET", url,
                                       params=params, json=payload, headers=headers,
                                       timeout=(15, 45))
            if response.status_code >= 400:
                try:
                    body = response.json()
                    detail = str(body.get("reason", body.get("detail", body.get("error", ""))))[:220]
                except ValueError:
                    detail = "Provider temporarily unavailable or request unsupported."
                raise DataError(f"{urlparse(url).hostname}: HTTP {response.status_code}. {detail}")
            data = response.json()
            if isinstance(data, dict) and data.get("error") is True:
                raise DataError(str(data.get("reason", "Provider rejected the request."))[:250])
            return data, utc_now()
        except requests.RequestException as exc:
            raise DataError(f"{urlparse(url).hostname}: connection failed ({type(exc).__name__}). Try again later.") from exc
        except ValueError as exc:
            raise DataError(f"{urlparse(url).hostname}: invalid JSON response.") from exc


def source_record(eid, provider, kind, retrieved, period, resolution, method, grid=None, url=None):
    return {"evidence_id": eid, "provider": provider, "evidence_type": kind,
            "retrieved_utc": retrieved, "period": period, "resolution": resolution,
            "method": method, "grid_coordinates": grid or "", "source_url": url or ""}


def result(name):
    return {"name": name, "tables": {}, "facts": [], "notes": [], "sources": [], "metrics": {}}


def fmt(value, digits=2):
    try:
        return f"{float(value):,.{digits}f}" if np.isfinite(float(value)) else "Unavailable"
    except (ValueError, TypeError):
        return "Unavailable"


def finite_stat(values, operation="mean"):
    values = pd.to_numeric(pd.Series(values), errors="coerce").dropna()
    if values.empty:
        return None
    return float(getattr(values, operation)())


def circle_geometry(lat, lon, radius_km):
    angles = np.linspace(0, 360, 97)
    xs, ys, _ = GEOD.fwd(np.full(97, lon), np.full(97, lat), angles, np.full(97, radius_km * 1000))
    return mapping(Polygon(zip(xs, ys)))


def normalize_geometry(raw):
    if raw.get("type") == "FeatureCollection":
        geoms = [shape(f["geometry"]) for f in raw.get("features", []) if f.get("geometry")]
        if not geoms:
            raise DataError("The GeoJSON does not contain a polygon.")
        geom = unary_union(geoms)
    elif raw.get("type") == "Feature":
        geom = shape(raw["geometry"])
    else:
        geom = shape(raw)
    if geom.geom_type not in ("Polygon", "MultiPolygon") or geom.is_empty or not geom.is_valid:
        raise DataError("Use a valid Polygon or MultiPolygon in WGS84 longitude/latitude.")
    west, south, east, north = geom.bounds
    if not (-180 <= west < east <= 180 and -80 <= south < north <= 80):
        raise DataError("Boundary coordinates must be WGS84, between 80°S and 80°N.")
    if east - west > 15 or north - south > 15:
        raise DataError("Select a local study area (under 15 degrees wide/high); antimeridian areas are not supported.")
    if len(json.dumps(mapping(geom))) > 400_000:
        raise DataError("Simplify this boundary before uploading (maximum geometry size: 400 KB).")
    return geom


def geometry_area(geom):
    parts = list(geom.geoms) if geom.geom_type == "MultiPolygon" else [geom]
    return sum(abs(GEOD.geometry_area_perimeter(orient(p, sign=1.0))[0]) for p in parts) / 1e6


def make_study(label, lat, lon, radius, start, end, custom=None):
    if not (-80 <= lat <= 80 and -180 <= lon <= 180):
        raise DataError("Enter a valid latitude and longitude.")
    if start > end:
        raise DataError("The start date must precede the end date.")
    if end > date.today():
        raise DataError("The study period is historical. Forecasts use a separate future window.")
    if (end - start).days > 3660:
        raise DataError("Limit the historical study to ten years per run. ")
    geom = normalize_geometry(custom or circle_geometry(lat, lon, radius))
    center = geom.centroid
    return {"label": str(label).strip()[:160] or "Selected study area", "lat": center.y,
            "lon": center.x, "geometry": mapping(geom), "bbox": list(geom.bounds),
            "area_km2": geometry_area(geom), "start": str(start), "end": str(end),
            "boundary": "Uploaded/drawn polygon" if custom else f"{radius:g} km radius"}


@st.cache_data(ttl=86400, max_entries=64, show_spinner=False)
def city_search(query):
    data, _ = request_json("https://geocoding-api.open-meteo.com/v1/search",
                           {"name": query, "count": 8, "language": "en", "format": "json"})
    return [{"label": ", ".join(str(r[k]) for k in ["name", "admin1", "country"] if r.get(k)),
             "lat": r["latitude"], "lon": r["longitude"]} for r in data.get("results", [])]


@st.cache_resource
def geocoder_limiter():
    return {"lock": threading.Lock(), "last": 0.0}


@st.cache_data(ttl=604800, max_entries=128, show_spinner=False)
def landmark_search(query):
    # Application-wide throttling and user-triggered searches only; no autocomplete.
    limiter = geocoder_limiter()
    with limiter["lock"]:
        wait = 1.1 - (time.monotonic() - limiter["last"])
        if wait > 0:
            time.sleep(wait)
        try:
            data, _ = request_json("https://nominatim.openstreetmap.org/search",
                                   {"q": query, "format": "jsonv2", "limit": 5})
        finally:
            limiter["last"] = time.monotonic()
    return [{"label": r["display_name"], "lat": float(r["lat"]), "lon": float(r["lon"])} for r in data]


def daily_frame(data):
    raw = data.get("daily", {})
    if not raw.get("time"):
        raise DataError("The provider returned no daily records for this period.")
    frame = pd.DataFrame(raw).rename(columns={"time": "date"})
    frame["date"] = pd.to_datetime(frame["date"])
    return frame


















@st.cache_data(ttl=3600, max_entries=16, show_spinner=False)
def satellite_catalogue(study, cloud_limit=40):
    # Bounded pagination. Full-month/date coverage is never implied by this subset.
    params = {"collections": SAT_COLLECTION, "bbox": ",".join(map(str, study["bbox"])),
              "datetime": f"{study['start']}T00:00:00Z/{study['end']}T23:59:59Z", "limit": 100,
              "sortby": "-properties.datetime"}
    url, features, truncated, stamp = "https://earth-search.aws.element84.com/v1/search", [], False, ""
    for page in range(2):
        data, stamp = request_json(url, params)
        features.extend(data.get("features", []))
        nxt = next((l for l in data.get("links", []) if l.get("rel") == "next"), None)
        if not nxt:
            break
        if page == 1:
            truncated = True
            break
        next_url = nxt["href"]
        if urlparse(next_url).hostname != "earth-search.aws.element84.com":
            break
        url, params = next_url, None
    features = [f for f in features if f.get("properties", {}).get("eo:cloud_cover", 100) <= cloud_limit]
    # Prefer the scene covering most of the AOI when multiple tiles/datetakes share a date.
    area = shape(study["geometry"])
    by_date = {}
    for f in features:
        day = f["properties"]["datetime"][:10]
        score = shape(f["geometry"]).intersection(area).area
        if day not in by_date or score > by_date[day][0]:
            by_date[day] = (score, f)
    return sorted([x[1] for x in by_date.values()], key=lambda f: f["properties"]["datetime"]), stamp, truncated


def spectral_index(a, b):
    a, b = np.asarray(a, dtype="float32"), np.asarray(b, dtype="float32")
    denom = a + b
    valid = np.isfinite(a) & np.isfinite(b) & (a >= 0) & (b >= 0) & (denom > 1e-6)
    out = np.full_like(a, np.nan, dtype="float32")
    np.divide(a - b, denom, out=out, where=valid)
    return out


def satellite_grid(study):
    import rasterio.transform
    zone = min(60, max(1, int((study["lon"] + 180) / 6) + 1))
    epsg = 32600 + zone if study["lat"] >= 0 else 32700 + zone
    project = Transformer.from_crs(4326, epsg, always_xy=True).transform
    geom = shape_transform(project, shape(study["geometry"]))
    x0, y0, x1, y1 = geom.bounds
    resolution = max(20, math.ceil(math.sqrt((x1 - x0) * (y1 - y0) / MAX_SAT_PIXELS) / 20) * 20)
    width, height = math.ceil((x1 - x0) / resolution), math.ceil((y1 - y0) / resolution)
    transform = rasterio.transform.from_origin(x0, y1, resolution, resolution)
    return epsg, transform, width, height, geom, resolution


def asset_for(item, names):
    for name in names:
        a = item.get("assets", {}).get(name)
        if a and a.get("href", "").lower().endswith((".tif", ".tiff")):
            host = urlparse(a["href"]).hostname or ""
            if urlparse(a["href"]).scheme != "https" or not host.endswith(".amazonaws.com"):
                raise DataError("Satellite asset host is outside the trusted Earth Search data hosts.")
            return a
    raise DataError(f"This satellite scene lacks a supported COG asset: {names[0]}.")


def read_satellite_asset(asset, grid, categorical=False):
    import rasterio
    from rasterio.enums import Resampling
    from rasterio.vrt import WarpedVRT
    epsg, transform, width, height, _, _ = grid
    band_meta = asset.get("raster:bands", [{}])[0]
    if not categorical and "scale" not in band_meta:
        raise DataError("Reflectance scale metadata is missing; the scene cannot be safely analysed.")
    ca_bundle = os.getenv("CURL_CA_BUNDLE") or os.getenv("REQUESTS_CA_BUNDLE") or os.getenv("SSL_CERT_FILE") or requests.certs.where()
    with rasterio.Env(GDAL_CURL_CA_BUNDLE=ca_bundle, CURL_CA_BUNDLE=ca_bundle, GDAL_DISABLE_READDIR_ON_OPEN="EMPTY_DIR", GDAL_HTTP_TIMEOUT="60",
                      GDAL_HTTP_CONNECTTIMEOUT="15", GDAL_HTTP_MAX_RETRY="1", GDAL_HTTP_RETRY_DELAY="1",
                      CPL_VSIL_CURL_ALLOWED_EXTENSIONS=".tif,.tiff", GDAL_CACHEMAX=64_000_000):
        with rasterio.open(asset["href"]) as src:
            with WarpedVRT(src, crs=f"EPSG:{epsg}", transform=transform, width=width, height=height,
                           src_nodata=src.nodata if src.nodata is not None else 0, nodata=-9999,
                           dtype="float32", resampling=Resampling.nearest if categorical else Resampling.bilinear) as vrt:
                arr = vrt.read(1, masked=True).filled(np.nan).astype("float32")
    if not categorical:
        # Apply STAC scale and offset exactly once; do not infer from processing baseline.
        arr = arr * float(band_meta["scale"]) + float(band_meta.get("offset", 0))
    return arr


def process_satellite_item(item, study, grid, water_threshold=0.0):
    from rasterio.features import geometry_mask
    epsg, transform, width, height, geom, resolution = grid
    inside = geometry_mask([mapping(geom)], (height, width), transform, invert=True)
    scl = read_satellite_asset(asset_for(item, ["scl", "SCL"]), grid, True)
    valid = inside & np.isin(scl, [4, 5, 6])
    if valid.sum() < 10:
        raise DataError("Fewer than ten clear, supported pixels remain inside the study area.")
    choices = {"blue": ["blue", "B02"], "green": ["green", "B03"], "red": ["red", "B04"],
               "nir": ["nir", "B08"], "rededge": ["rededge1", "B05"], "swir": ["swir16", "B11"]}
    assets = {name: asset_for(item, keys) for name, keys in choices.items()}
    radiometry = {name: asset.get("raster:bands", [{}])[0] for name, asset in assets.items()}
    # Independent COG windows; keep concurrency low for shared Streamlit instances.
    with ThreadPoolExecutor(max_workers=3) as pool:
        futures = {name: pool.submit(read_satellite_asset, asset, grid) for name, asset in assets.items()}
        bands = {name: future.result() for name, future in futures.items()}
    for band in bands.values():
        valid &= np.isfinite(band) & (band >= 0)
    ndvi = spectral_index(bands["nir"], bands["red"])
    ndwi = spectral_index(bands["green"], bands["nir"])
    mndwi = spectral_index(bands["green"], bands["swir"])
    water = valid & (scl == 6) & ((ndwi > water_threshold) | (mndwi > water_threshold)) & (ndvi < 0.3)
    ndci = spectral_index(bands["rededge"], bands["red"])
    ndci[~water] = np.nan
    for arr in [ndvi, ndwi, mndwi]:
        arr[~valid] = np.nan
    redwater = np.where(water, bands["red"], np.nan).astype("float32")
    rgb = np.stack([bands["red"], bands["green"], bands["blue"]], axis=-1)
    rgb = np.clip(np.nan_to_num(rgb) / 0.3, 0, 1) ** (1 / 1.8)
    rgb[~valid] = 0
    arrays = {"NDWI": ndwi, "MNDWI": mndwi, "NDCI": ndci,
              "Water red reflectance": redwater, "Water mask": np.where(valid, water.astype(float), np.nan).astype("float32")}
    pixel_km2 = resolution ** 2 / 1e6
    summary = {"scene_id": item["id"], "date": item["properties"]["datetime"],
        "scene_cloud_percent": item["properties"].get("eo:cloud_cover"),
        "valid_aoi_percent": 100 * valid.sum() / max(1, inside.sum()),
        "screened_water_km2": float(water.sum() * pixel_km2), "water_pixels": int(water.sum()),
        "median_water_ndci": float(np.nanmedian(ndci)) if np.isfinite(ndci).any() else np.nan,
        "median_water_red_reflectance": float(np.nanmedian(redwater)) if np.isfinite(redwater).any() else np.nan,
        "grid_resolution_m": resolution, "collection": item.get("collection", SAT_COLLECTION)}
    return {"summary": summary, "arrays": arrays, "rgb": rgb.astype("float32"), "valid": valid,
            "inside": inside, "water": water, "epsg": epsg, "transform": tuple(transform),
            "resolution": resolution, "radiometry": radiometry}


@st.cache_data(ttl=3600, max_entries=2, show_spinner=False)
def satellite_module(study, scene_count=3, cloud_limit=40, water_threshold=0.0):
    if study["area_km2"] > MAX_SAT_KM2:
        raise DataError(f"Satellite processing is limited to {MAX_SAT_KM2:g} km² per run. Draw a smaller study area.")
    if (date.fromisoformat(study["end"]) - date.fromisoformat(study["start"])).days > 1096:
        raise DataError("Use a satellite study period of three years or less per run.")
    items, stamp, truncated = satellite_catalogue(study, cloud_limit)
    if not items:
        raise DataError("No Sentinel-2 scenes pass the scene-cloud filter. Expand the dates or cloud limit.")
    indices = ([len(items) - 1] if scene_count == 1 else
               np.unique(np.linspace(0, len(items) - 1, min(scene_count, len(items))).round().astype(int)))
    grid, outputs, failures = satellite_grid(study), [], []
    for idx in indices:
        item = items[idx]
        try:
            outputs.append(process_satellite_item(item, study, grid, water_threshold))
        except Exception as exc:
            failures.append(f"{item['id']}: {str(exc)[:220]}")
    if not outputs:
        raise DataError("No selected scene could be processed. " + "; ".join(failures)[:650])
    out = result("Satellite")
    out["tables"]["Satellite scene statistics"] = pd.DataFrame([o["summary"] for o in outputs])
    out["raster"] = outputs[-1]
    last = outputs[-1]["summary"]
    out["metrics"]["Latest screened water (km²)"] = last["screened_water_km2"]
    out["metrics"]["Latest usable AOI (%)"] = last["valid_aoi_percent"]
    out["facts"].append(f"[S1] Latest processed scene: {last['date'][:10]}; {last['valid_aoi_percent']:.1f}% usable AOI coverage; {last['screened_water_km2']:.3f} km² of screened water in the valid footprint.")
    if last["water_pixels"] == 0:
        out["facts"].append("[S1] No pixels passed the water screen in the latest processed scene. Water-quality indices and sampling candidates are unavailable; this does not establish that water is absent or clean.")
    else:
        out["facts"].append(f"[S1] Median water NDCI: {fmt(last['median_water_ndci'], 3)}. NDCI is a dimensionless screening indicator, not chlorophyll concentration or bloom confirmation.")
    if last["valid_aoi_percent"] < 50:
        out["notes"].append(f"Only {last['valid_aoi_percent']:.1f}% of the latest scene's study area passed the quality mask. Conditions in the excluded area are unknown; inspect an earlier clearer scene before drawing conclusions.")
    if len(outputs) >= 2:
        first, last_raster = outputs[0], outputs[-1]
        common = first["valid"] & last_raster["valid"]
        scale = first["resolution"] ** 2 / 1e6
        comparison = {"first_date": first["summary"]["date"], "last_date": last_raster["summary"]["date"],
            "common_valid_km2": float(common.sum() * scale),
            "first_water_km2": float((first["water"] & common).sum() * scale),
            "last_water_km2": float((last_raster["water"] & common).sum() * scale)}
        comparison["water_change_km2"] = comparison["last_water_km2"] - comparison["first_water_km2"]
        out["tables"]["Common footprint comparison"] = pd.DataFrame([comparison])
        if comparison["first_water_km2"] == 0 and comparison["last_water_km2"] == 0:
            out["notes"].append("Neither comparison date detected water in the common valid footprint. A zero difference is not evidence of stable river extent; water-area change is inconclusive.")
        elif common.sum() >= 10:
            out["facts"].append(f"[S1] Screened water changed by {comparison['water_change_km2']:+.3f} km² within the {comparison['common_valid_km2']:.3f} km² footprint valid on both comparison dates.")
        else:
            out["notes"].append("The common valid footprint is too small for a meaningful change assessment.")
    # Relative within-scene NDCI hotspots support field sampling, not a universal risk threshold.
    ndci = out["raster"]["arrays"]["NDCI"]
    yy, xx = np.where(np.isfinite(ndci))
    if len(xx):
        from affine import Affine
        aff = Affine(*out["raster"]["transform"][:6])
        inv = Transformer.from_crs(out["raster"]["epsg"], 4326, always_xy=True)
        order, chosen, rows = np.argsort(ndci[yy, xx])[::-1], [], []
        for i in order:
            x, y = aff * (int(xx[i]) + 0.5, int(yy[i]) + 0.5)
            if any(math.hypot(x - px0, y - py0) < 200 for px0, py0 in chosen):
                continue
            lon, lat = inv.transform(x, y)
            chosen.append((x, y))
            rows.append({"priority_rank": len(rows) + 1, "latitude": lat, "longitude": lon,
                         "ndci": float(ndci[yy[i], xx[i]]), "date": last["date"], "basis": "Relative NDCI rank; unverified sampling candidate"})
            if len(rows) == 10:
                break
        out["tables"]["Sampling candidates"] = pd.DataFrame(rows)
    out["sources"].append(source_record("S1", "Copernicus Sentinel-2 Collection 1 L2A / Earth Search", "Satellite-derived screening", stamp,
        f"{outputs[0]['summary']['date']} to {last['date']}", f"Common UTM grid: {grid[-1]} m; native source bands 10/20 m",
        "STAC radiometric scaling; bilinear reflectance; nearest-neighbour SCL. Only SCL 4/5/6; screened water additionally SCL 6, NDVI<0.3 and NDWI or MNDWI above the chosen threshold.", url=SOURCES["Satellite"]))
    out["notes"].extend([
        "Uses Earth Search sentinel-2-c1-l2a only. Catalogue gaps remain unavailable; the legacy collection is not used as a calibration fallback.",
        "Selected scenes only, not a continuous monthly time series. Whole-scene cloud cover differs from usable coverage inside the AOI.",
        "NDCI uses Sentinel-2 B5/B4. Red reflectance is an uncalibrated sediment-related optical proxy, not NTU, TSS or nutrient concentration.",
        "L2A land surface-reflectance correction has limitations over water. Shoreline mixing, glint, sediment and aquatic vegetation can confound indices; field validation is required.",
        "Water-area change is compared only in a common clear footprint. Very narrow rivers may be unresolved. Sampling candidates are relative ranks, not confirmed pollution hotspots."])
    if truncated:
        out["notes"].append("The catalogue search reached its 200-item cap; selected scenes cover only the retrieved subset. Narrow the dates for fuller coverage.")
    out["notes"].extend(failures)
    return out






def all_sources(run):
    return [s for r in run["results"].values() for s in r.get("sources", [])]


def all_tables(run):
    return [(module, name, frame) for module, r in run["results"].items() for name, frame in r.get("tables", {}).items()]




def interactive_chart(spec):
    df = spec["df"].head(12) if spec["kind"] == "horizontal" else spec["df"]
    if spec["kind"] == "horizontal":
        fig = px.bar(df, y=spec["x"], x=spec["ys"][0], orientation="h", color_discrete_sequence=COLORS)
    elif spec["kind"] == "bar":
        fig = px.bar(df, x=spec["x"], y=spec["ys"], color_discrete_sequence=COLORS, barmode="group")
    elif spec["kind"] == "scatter":
        fig = px.scatter(df, x=spec["x"], y=spec["ys"], color_discrete_sequence=COLORS)
    else:
        fig = px.line(df, x=spec["x"], y=spec["ys"], color_discrete_sequence=COLORS, markers=len(df) < 35)
    fig.update_layout(template="plotly_dark", paper_bgcolor="rgba(0,0,0,0)", plot_bgcolor="rgba(0,0,0,0)",
                      font=dict(color="#DFE7F1"), height=345, title=None, margin=dict(l=12, r=15, t=15, b=20),
                      yaxis_title=spec["units"], xaxis_title=None, legend_title=None,
                      legend=dict(orientation="h", y=-0.25, x=0))
    fig.update_xaxes(showgrid=False)
    if spec.get("date_range"):
        start, end = pd.to_datetime(spec["date_range"])
        fig.update_xaxes(range=[start-pd.Timedelta(hours=12), end+pd.Timedelta(days=1)])
    fig.update_yaxes(gridcolor="rgba(170,180,210,.12)")
    return fig


def figure_png(fig):
    buffer = io.BytesIO()
    fig.savefig(buffer, format="png", dpi=300, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return buffer.getvalue()


def plot_static(spec):
    df = spec["df"].head(12).iloc[::-1] if spec["kind"] == "horizontal" else spec["df"]
    fig, ax = plt.subplots(figsize=(9.6, 4.3))
    for i, y in enumerate(spec["ys"]):
        color = ["#087F8C", "#7552AD", "#BA7900"][i % 3]
        if spec["kind"] == "horizontal":
            ax.barh(df[spec["x"]].fillna("Unspecified").astype(str), df[y], color=color)
        elif spec["kind"] == "bar":
            ax.bar(df[spec["x"]], df[y], color=color, label=y, width=.7)
        elif spec["kind"] == "scatter":
            ax.scatter(df[spec["x"]], df[y], color=color, s=22, alpha=.8, label=y)
        else:
            xx = pd.to_datetime(df[spec["x"]]) if spec["x"] == "date" else df[spec["x"]]
            ax.plot(xx, df[y], color=color, lw=1.6, marker="o" if len(df) < 35 else None, markersize=3, label=y)
    ax.set_title(spec["title"], loc="left", fontsize=11, fontweight="bold", pad=15)
    ax.set_ylabel(spec["units"])
    if spec.get("date_range"):
        start, end = pd.to_datetime(spec["date_range"])
        ax.set_xlim(start-pd.Timedelta(hours=12), end+pd.Timedelta(days=1))
    ax.grid(axis="y", alpha=.18)
    ax.spines[["top", "right"]].set_visible(False)
    if len(spec["ys"]) > 1:
        ax.legend(fontsize=7)
    if spec["kind"] != "horizontal":
        fig.autofmt_xdate()
    fig.text(.12, -.015, f"Evidence: {spec['evidence']} | HydroScope Water Research", fontsize=8, color="#526174")
    return figure_png(fig)


def draw_boundary(ax, study):
    geom = shape(study["geometry"])
    for part in list(geom.geoms) if geom.geom_type == "MultiPolygon" else [geom]:
        x, y = part.exterior.xy
        ax.fill(x, y, color="#D6F4EC", alpha=.55)
        ax.plot(x, y, color="#087F8C", lw=1.5, label="Study boundary")
        for hole in part.interiors:
            hx, hy = hole.xy
            ax.fill(hx, hy, color="white")




RASTER_STYLES = {"NDWI": ("BrBG", -1, 1),
                 "MNDWI": ("BrBG", -1, 1), "NDCI": ("viridis", -.3, .6),
                 "Water red reflectance": ("YlOrBr", 0, .15), "Water mask": ("Blues", 0, 1)}


def raster_png(raster, name):
    from affine import Affine
    from rasterio.transform import array_bounds
    t = Affine(*raster["transform"][:6])
    h, w = raster["valid"].shape
    west, south, east, north = array_bounds(h, w, t)
    fig, ax = plt.subplots(figsize=(8.8, 5.5))
    if name == "True colour":
        rgba = np.dstack([raster["rgb"], raster["valid"].astype(float)])
        ax.imshow(rgba, extent=[west, east, south, north])
    elif not np.isfinite(raster["arrays"][name]).any():
        ax.set_xlim(west, east)
        ax.set_ylim(south, north)
        ax.text(.5, .5, "No usable water-index pixels\nThis indicator is unavailable; no water-quality conclusion.",
                transform=ax.transAxes, ha="center", va="center", color="#334155", fontsize=12,
                bbox={"facecolor": "#F1F5F9", "edgecolor": "#CBD5E1", "pad": 14})
    else:
        cmap, lo, hi = RASTER_STYLES[name]
        im = ax.imshow(raster["arrays"][name], extent=[west, east, south, north], cmap=cmap, vmin=lo, vmax=hi)
        fig.colorbar(im, ax=ax, shrink=.8, label="Reflectance" if "reflectance" in name else "Index / class", extend="both")
    ax.set_title(f"{name} | {raster['summary']['date'][:10]}", loc="left", fontweight="bold", fontsize=12)
    ax.set_xlabel(f"Easting (m), EPSG:{raster['epsg']}")
    ax.set_ylabel("Northing (m)")
    ax.ticklabel_format(style="plain", useOffset=False)
    from matplotlib.ticker import MaxNLocator
    ax.xaxis.set_major_locator(MaxNLocator(5))
    ax.yaxis.set_major_locator(MaxNLocator(5))
    ax.annotate("N", xy=(.94, .94), xytext=(.94, .82), xycoords="axes fraction", ha="center", color="#37284B", arrowprops=dict(arrowstyle="->", color="#37284B"))
    # Scale bar length is derived from the projected grid.
    length = max(raster["resolution"], round((east-west)/4 / 100) * 100)
    x0, y0 = west + (east-west)*.07, south + (north-south)*.08
    ax.plot([x0, x0+length], [y0, y0], lw=4, color="#F4C97A")
    ax.text(x0, y0+(north-south)*.035, f"{length:g} m", color="#66461A", fontsize=8)
    fig.text(.12, .005, f"S1: Copernicus Sentinel-2 / Earth Search | grid {raster['resolution']} m | blank = masked/no data", fontsize=8)
    return figure_png(fig)


def safe_frame(frame):
    """Prevent spreadsheet formula execution when exporting provider/user text."""
    output = frame.copy()
    for col in output.select_dtypes(include=["object", "string"]).columns:
        output[col] = output[col].map(lambda v: "'" + v if isinstance(v, str) and v.lstrip().startswith(("=", "+", "-", "@")) else v)
    for col in output.columns:
        if isinstance(output[col].dtype, pd.DatetimeTZDtype):
            output[col] = output[col].dt.tz_convert("UTC").dt.tz_localize(None)
    return output


def plain_metadata(run):
    def clean(value):
        if isinstance(value, dict):
            return {k: clean(v) for k, v in value.items()}
        if isinstance(value, (list, tuple)):
            return [clean(v) for v in value]
        if isinstance(value, np.generic):
            return clean(value.item())
        if isinstance(value, float) and not math.isfinite(value):
            return None
        return value
    metadata = {k: v for k, v in run.items() if k != "results"} | {
        "results": {name: {k: v for k, v in r.items() if k not in ["tables", "raster"]} for name, r in run["results"].items()},
        "table_inventory": [{"module": m, "table": t, "rows": len(f), "columns": list(f.columns)} for m, t, f in all_tables(run)]}
    raster = run["results"].get("Satellite", {}).get("raster")
    if raster:
        metadata["satellite_raster"] = {k: raster[k] for k in ["summary", "epsg", "transform", "resolution", "radiometry"]}
        metadata["satellite_raster"]["bands"] = list(raster["arrays"])
    return clean(metadata)


def geotiff_bytes(raster):
    from affine import Affine
    from rasterio.io import MemoryFile
    names = list(raster["arrays"])
    h, w = raster["valid"].shape
    with MemoryFile() as mem:
        with mem.open(driver="GTiff", width=w, height=h, count=len(names), dtype="float32", crs=f"EPSG:{raster['epsg']}",
                      transform=Affine(*raster["transform"][:6]), nodata=-9999.0, compress="deflate") as dst:
            for i, name in enumerate(names, 1):
                dst.write(np.where(np.isfinite(raster["arrays"][name]), raster["arrays"][name], -9999).astype("float32"), i)
                dst.set_band_description(i, name)
            dst.update_tags(scene=raster["summary"]["scene_id"], acquisition=raster["summary"]["date"],
                source="Copernicus Sentinel-2 L2A / Earth Search", product="Uncalibrated screening indices; not water-safety measurements")
        return mem.read()


def feature_collection(run):
    features = [{"type": "Feature", "geometry": run["study"]["geometry"], "properties": {"layer": "Study boundary", "name": run["study"]["label"]}}]
    for module, name, frame in all_tables(run):
        if name not in ["Sampling candidates", "Included field observations"] or frame.empty:
            continue
        for row in json.loads(frame.to_json(orient="records", date_format="iso")):
            if row.get("latitude") is None or row.get("longitude") is None:
                continue
            features.append({"type": "Feature", "geometry": {"type": "Point", "coordinates": [row["longitude"], row["latitude"]]},
                             "properties": {"layer": name, **row}})
    return {"type": "FeatureCollection", "features": features}




def current_interpretation(run):
    review = run.get("interpretation")
    if review:
        from evidence import interpretation_is_current
        if not interpretation_is_current(run, review):
            raise DataError("The AI review is incomplete or belongs to different evidence. Generate a new interpretation before including it.")
    return review


def build_html_report(run, figures):
    e = lambda x: html.escape(str(x))
    study = run["study"]
    parts = ["<!doctype html><html lang='en'><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'>",
        "<title>HydroScope Water Research report</title><style>body{font:15px/1.65 system-ui;color:#172638;background:#f2f5f8;margin:0}main{max-width:1050px;margin:auto;padding:38px;background:white}header{background:#12132c;color:white;padding:35px;border-top:6px solid #5fe1c3;border-radius:15px}h1{font-size:32px}h2{color:#087f8c;margin-top:35px}h3{color:#5a437b}table{border-collapse:collapse;width:100%;font-size:12px;display:block;overflow:auto}td,th{border:1px solid #dfe5ec;padding:7px;text-align:left}th{background:#edf5f4}figure{margin:28px 0;break-inside:avoid}img{width:100%;max-width:950px}figcaption,.note{color:#58677b;font-size:12px}.warning{background:#fff4dc;padding:16px}.source{overflow-wrap:anywhere}li{margin:7px 0}@media print{body{background:white}main{padding:0}header{-webkit-print-color-adjust:exact}h2{break-after:avoid}}</style><main>",
        f"<header><p>HYDROSCOPE WATER RESEARCH · WATER RESEARCH</p><h1>{e(study['label'])}</h1><p>{e(study['start'])} — {e(study['end'])}</p><p>Run {e(run['id'])} · generated {e(run['created_utc'])}</p></header>",
        f"<p>Study area: {study['area_km2']:.3f} km² · {e(study['boundary'])} · centroid {study['lat']:.5f}, {study['lon']:.5f}. Forecast windows are stated separately.</p>",
        "<div class='warning'>Research and screening report. Satellite indicators require field validation. River forecasts require local river/gauge checks. Marine model output is not suitable for navigation or water-safety decisions.</div>",
        "<h2>Coverage and unavailable evidence</h2><ul>" + "".join(f"<li>{e(n)}</li>" for n in report_scope_notes(run)) + "</ul>",
        "<h2>Executive findings</h2><ul>"]
    for r in run["results"].values():
        parts.extend(f"<li>{e(f)}</li>" for f in r["facts"])
    parts.append("</ul>")
    if not run["results"]:
        parts.append("<p>No data module completed. This report documents the unsuccessful retrieval; it does not assess environmental conditions.</p>")
    if run["errors"]:
        parts.append("<h2>Unavailable modules</h2><ul>" + "".join(f"<li>{e(k)}: {e(v)}</li>" for k,v in run["errors"].items()) + "</ul>")
    review = current_interpretation(run)
    if review:
        parts.append("<h2>Optional AI interpretation</h2><p class='note'>AI-generated interpretation. This is a single-model explanation of saved calculations, not independent scientific validation. Verify the cited evidence.</p>")
        parts.append(f"<p>Model: {e(review['model'])} · Generated: {e(review['generated_utc'])}</p>")
        parts.append("<div style='white-space:pre-wrap;overflow-wrap:anywhere'>" + e(review["answer"]) + "</div>")
    parts.append("<h2>Maps and figures</h2>")
    for title, content in figures:
        parts.append(f"<figure><img alt='{e(title)}' src='data:image/png;base64,{base64.b64encode(content).decode()}'><figcaption>{e(title)}</figcaption></figure>")
    for module, r in run["results"].items():
        parts.append(f"<h2>{e(module)}</h2>")
        for name, frame in r["tables"].items():
            parts.append(f"<h3>{e(name)}</h3><p class='note'>Showing up to 20 of {len(frame)} rows. Full records are supplied in the data exports.</p>")
            parts.append(frame.head(20).to_html(index=False, escape=True, float_format=lambda x: f"{x:.4g}", na_rep="Unavailable"))
        parts.append("<ul>" + "".join(f"<li>{e(n)}</li>" for n in r["notes"]) + "</ul>")
    parts.append("<h2>Recommended follow-up</h2><ul><li>Verify the study boundary and the spatial support of each dataset.</li><li>Sample candidate water locations and compare measurements with satellite acquisition dates.</li><li>Confirm river-cell identity and local gauge thresholds before interpreting discharge forecasts.</li><li>Retain missing-data and uncertainty notes when sharing results.</li></ul><h2>Sources and reproducibility</h2>")
    for s in all_sources(run):
        parts.append("<p class='source'>" + "<br>".join(f"<b>{e(k.replace('_',' '))}:</b> {e(v)}" for k,v in s.items()) + "</p>")
    parts.append(f"<p class='note'>HydroScope Water Research v{VERSION}. Accompanying metadata.json retains the selected boundary and analysis settings. No sample data were substituted for failed sources.</p></main></html>")
    return "".join(parts).encode("utf-8")


def build_pdf_report(run, figures):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image, Table, TableStyle, PageBreak, KeepTogether
    from matplotlib.font_manager import findfont
    pdfmetrics.registerFont(TTFont("EcoSans", findfont("DejaVu Sans")))
    pdfmetrics.registerFont(TTFont("EcoSansBold", findfont("DejaVu Sans:weight=bold")))
    pdfmetrics.registerFontFamily("EcoSans", normal="EcoSans", bold="EcoSansBold", italic="EcoSans", boldItalic="EcoSansBold")
    styles = getSampleStyleSheet()
    for style in styles.byName.values():
        style.fontName = "EcoSans"
    styles["BodyText"].fontSize, styles["BodyText"].leading = 8.5, 13
    styles["Title"].fontName = styles["Heading1"].fontName = "EcoSansBold"
    styles["Heading1"].textColor = colors.HexColor("#087F8C")
    small = ParagraphStyle("EcoSmall", parent=styles["BodyText"], fontSize=6.8, leading=9)
    p = lambda text, style="BodyText": Paragraph(html.escape(str(text)), styles[style])
    output, flow = io.BytesIO(), []
    doc = SimpleDocTemplate(output, pagesize=A4, leftMargin=1.6*cm, rightMargin=1.6*cm,
                            topMargin=1.8*cm, bottomMargin=1.8*cm, title=f"HydroScope Water Research — {run['study']['label']}", author="HydroScope Water Research")
    flow += [p("HYDROSCOPE WATER RESEARCH", "Title"), p(run["study"]["label"], "Heading1"),
             p(f"Historical study: {run['study']['start']} to {run['study']['end']}"),
             p(f"Run {run['id']} | Generated {run['created_utc']}"), Spacer(1,.4*cm),
             p(f"Boundary area {run['study']['area_km2']:.3f} km². Forecast windows are separate. Research/screening output; validate water indicators and river forecasts locally. Marine output is not suitable for navigation or water-safety decisions."),
             p("Coverage and unavailable evidence", "Heading1")]
    flow.extend(p("• " + n) for n in report_scope_notes(run))
    flow.append(p("Executive findings", "Heading1"))
    if not run["results"]:
        flow.append(p("No data module completed. Environmental conditions cannot be assessed from this run."))
    for r in run["results"].values():
        flow.extend(p("• " + f) for f in r["facts"])
    for module, error in run["errors"].items():
        flow.append(p(f"Unavailable — {module}: {error}"))
    review = current_interpretation(run)
    if review:
        flow += [p("Optional AI interpretation", "Heading1"),
                 p("AI-generated interpretation; this is a single-model explanation, not independent scientific validation. Verify cited evidence."),
                 p(f"Model: {review['model']} | Generated: {review['generated_utc']}")]
        for line in review["answer"].splitlines():
            if line.strip():
                heading = line.startswith("#")
                flow.append(p(line.lstrip("# ") if heading else line.replace("**", ""), "Heading2" if heading else "BodyText"))
    flow.append(PageBreak())
    from PIL import Image as PILImage
    for title, data in figures:
        with PILImage.open(io.BytesIO(data)) as im:
            w,h = im.size
        display_w = min(doc.width, doc.width * 1.0)
        display_h = display_w * h / w
        if display_h > 14*cm:
            display_w *= 14*cm / display_h
            display_h = 14*cm
        flow.append(KeepTogether([p(title, "Heading2"), Image(io.BytesIO(data), width=display_w, height=display_h), Spacer(1,.3*cm)]))
    for module, r in run["results"].items():
        flow.append(p(module, "Heading1"))
        for name, frame in r["tables"].items():
            if frame.empty:
                flow.append(p(f"{name}: no returned records."))
                continue
            # The PDF gives a legible preview; the workbook/CSV contains every column and row.
            cols = list(frame.columns[:6])
            preview = frame[cols].head(10).copy()
            text = lambda v: "Unavailable" if v is None or (not isinstance(v, str) and pd.isna(v)) else str(v)[:110]
            cells = [[Paragraph(html.escape(str(c)), small) for c in cols]]
            cells += [[Paragraph(html.escape(text(v)), small) for v in row] for row in preview.itertuples(index=False, name=None)]
            table = Table(cells, colWidths=[doc.width/len(cols)]*len(cols), repeatRows=1, hAlign="LEFT")
            table.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,0), colors.HexColor("#DEF1ED")),
                ("GRID", (0,0), (-1,-1), .3, colors.HexColor("#D8E2EA")), ("VALIGN", (0,0), (-1,-1), "TOP"),
                ("TOPPADDING", (0,0), (-1,-1), 5), ("BOTTOMPADDING", (0,0), (-1,-1), 5)]))
            flow += [p(name, "Heading2"), p(f"Preview: up to 10 of {len(frame)} rows and 6 columns. Use the workbook/CSV for complete data."), table, Spacer(1,.3*cm)]
        flow.extend(p("• " + n) for n in r["notes"])
    flow.append(p("Sources and reproducibility", "Heading1"))
    for s in all_sources(run):
        flow += [p(f"{s['evidence_id']} — {s['provider']}", "Heading2")]
        flow.extend(p(f"{k.replace('_',' ')}: {v}") for k,v in s.items() if k not in ["evidence_id", "provider"])
    def footer(canvas, document):
        canvas.setFont("EcoSans", 7)
        canvas.setFillColor(colors.HexColor("#607085"))
        canvas.drawString(1.6*cm, 1*cm, f"HydroScope Water Research | {run['id']} | Research and screening")
        canvas.drawRightString(A4[0]-1.6*cm, 1*cm, str(document.page))
    doc.build(flow, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


def build_exports(run, additional_figures=None):
    review = current_interpretation(run)
    figures = [("Study boundary and sampling map", point_map_png(run))]
    satellite = run["results"].get("Satellite", {}).get("raster")
    if satellite:
        for name in ["True colour", *RASTER_STYLES.keys()]:
            if name != "True colour" and not np.isfinite(satellite["arrays"][name]).any():
                continue
            figures.append((f"Satellite — {name}", raster_png(satellite, name)))
    figures.extend((s["title"], plot_static(s)) for s in chart_specs(run))
    figures.extend(additional_figures or [])
    metadata = json.dumps(plain_metadata(run), indent=2, ensure_ascii=False, default=str).encode()
    html_bytes = build_html_report(run, figures)
    pdf_bytes = build_pdf_report(run, figures)
    workbook = io.BytesIO()
    with pd.ExcelWriter(workbook, engine="openpyxl") as writer:
        index = []
        for i, (module, title, frame) in enumerate(all_tables(run), 1):
            sheet = f"{i:02d}_" + re.sub(r"[\[\]:*?/\\]", "", title)[:27]
            safe_frame(frame).to_excel(writer, sheet_name=sheet, index=False)
            index.append({"module": module, "table": title, "sheet": sheet, "rows": len(frame)})
        pd.DataFrame(index).to_excel(writer, sheet_name="INDEX", index=False)
        safe_frame(pd.DataFrame(all_sources(run))).to_excel(writer, sheet_name="SOURCES", index=False)
        pd.DataFrame([{"module": k, "error": v} for k,v in run["errors"].items()]).to_excel(writer, sheet_name="UNAVAILABLE", index=False)
        safe_frame(pd.DataFrame([{"run_id": run["id"], "generated_utc": run["created_utc"], "location": run["study"]["label"],
                       "area_km2": run["study"]["area_km2"], "start": run["study"]["start"], "end": run["study"]["end"]}])).to_excel(writer, sheet_name="STUDY", index=False)
        for ws in writer.book.worksheets:
            ws.freeze_panes = "A2"
            if ws.max_row > 1:
                ws.auto_filter.ref = ws.dimensions
            from openpyxl.styles import Font, PatternFill
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="174C59")
            for column in ws.columns:
                ws.column_dimensions[column[0].column_letter].width = min(45, max(15, len(str(column[0].value or ""))+3))
    bundle = io.BytesIO()
    with zipfile.ZipFile(bundle, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("report.html", html_bytes)
        z.writestr("report.pdf", pdf_bytes)
        z.writestr("analysis.xlsx", workbook.getvalue())
        z.writestr("metadata.json", metadata)
        if review:
            z.writestr("interpretation/results_interpretation.md",review["answer"])
            z.writestr("interpretation/provenance.json",json.dumps(review,indent=2,ensure_ascii=False))
        z.writestr("maps/study_and_observations.geojson", json.dumps(feature_collection(run), ensure_ascii=False))
        if satellite:
            z.writestr("maps/satellite_indices.tif", geotiff_bytes(satellite))
        for i, (module, title, frame) in enumerate(all_tables(run), 1):
            slug = re.sub(r"[^a-zA-Z0-9]+", "_", title).strip("_").lower()
            z.writestr(f"data/{i:02d}_{slug}.csv", safe_frame(frame).to_csv(index=False).encode("utf-8-sig"))
        for i, (title, data) in enumerate(figures, 1):
            z.writestr(f"figures/{i:02d}_{re.sub(r'[^a-zA-Z0-9]+', '_', title)[:70]}.png", data)
        z.writestr("READ_ME.txt", "HydroScope Water Research research output. See report limitations and metadata.json. CSV/XLSX contain full returned records. PDF/HTML tables are previews. Satellite GeoTIFF uses its recorded UTM CRS and -9999 nodata. Field observations remain unverified.\n")
    return {"pdf": pdf_bytes, "html": html_bytes, "xlsx": workbook.getvalue(), "zip": bundle.getvalue(), "metadata": metadata}




def execute_analysis(study, options, progress=None):
    from water_data import available_modules, river_module, marine_module
    run = {"id": hashlib.sha256((json.dumps(study, sort_keys=True) + str(time.time_ns())).encode()).hexdigest()[:12],
           "created_utc": utc_now(), "version": VERSION, "study": study, "options": options, "results": {}, "errors": {}}
    allowed = available_modules(study["waterbody_type"])
    adapters = {
        "Satellite": lambda: satellite_module(study, options.get("scene_count", 3), options.get("cloud_limit", 40), options.get("water_threshold", 0)),
        "River outlook": lambda: river_module(study, options.get("flow_threshold", 0)),
        "Marine outlook": lambda: marine_module(study),
    }
    for name in options.get("modules", []):
        if progress:
            progress("Retrieving and checking " + name.lower() + "…")
        try:
            if name not in allowed:
                raise DataError("This analysis is not available for the selected waterbody type.")
            run["results"][name] = adapters[name]()
        except Exception as exc:
            run["errors"][name] = str(exc)[:600] if isinstance(exc, DataError) else f"{type(exc).__name__}: provider processing could not complete. No replacement data were created."
    return run


def chart_specs(run):
    specs = []
    def add(module, table, title, x, ys, units, eid, kind="line"):
        f = run.get("results", {}).get(module, {}).get("tables", {}).get(table)
        if f is None or f.empty or x not in f:
            return
        columns = [y for y in ys if y in f and pd.to_numeric(f[y], errors="coerce").notna().any()]
        if columns:
            frame = f.copy()
            if x == "date":
                frame[x] = pd.to_datetime(frame[x], utc=True, errors="coerce").dt.tz_localize(None)
            specs.append({"title": title, "module": module, "table": table, "df": frame, "x": x,
                          "ys": columns, "units": units, "kind": kind, "evidence": eid})
    add("River outlook", "Discharge forecast", "Modelled discharge and ensemble quartiles", "date", ["river_discharge", "river_discharge_p25", "river_discharge_p75"], "m³/s", "F1")
    add("River outlook", "Historical modelled discharge", "Historical modelled discharge — not gauge measurements", "date", ["river_discharge"], "m³/s", "F2")
    for col, label, unit in [("sea_surface_temperature", "Modelled sea-surface temperature", "°C"),
                             ("wave_height", "Significant wave height", "m"),
                             ("ocean_current_velocity", "Modelled ocean current speed", "km/h"),
                             ("sea_level_height_msl", "Modelled sea level relative to global MSL", "m")]:
        add("Marine outlook", "Marine hourly forecast", label, "date", [col], unit, "M1")
    add("Satellite", "Satellite scene statistics", "Screened water area within each clear footprint", "date", ["screened_water_km2"], "km²", "S1")
    add("Satellite", "Satellite scene statistics", "Median water NDCI on selected dates", "date", ["median_water_ndci"], "Dimensionless index", "S1")
    for col, unit in [("chlorophyll_ug_l", "µg/L"), ("turbidity_ntu", "NTU"), ("secchi_m", "m"),
                      ("temperature_c", "°C"), ("dissolved_oxygen_mg_l", "mg/L"), ("salinity_psu", "PSU")]:
        add("Field observations", "Included field observations", "Field measurements: " + col, "date", [col], unit, "U1", "scatter")
    return specs


def point_map_png(run, module=None):
    fig, ax = plt.subplots(figsize=(9.2, 5.4))
    draw_boundary(ax, run["study"])
    for _, name, frame in all_tables(run):
        if name not in ("Sampling candidates", "Included field observations") or frame.empty:
            continue
        f = frame.dropna(subset=["latitude", "longitude"])
        if f.empty:
            continue
        sizes = 38
        label = name
        if name == "Included field observations" and "chlorophyll_ug_l" in f:
            # Draw non-missing measurements with area proportional to magnitude, with display bounds.
            sizes = np.clip(pd.to_numeric(f.chlorophyll_ug_l, errors="coerce").fillna(0).to_numpy() * 6, 18, 320)
            label += " (bubble size: chlorophyll; capped)"
        ax.scatter(f.longitude, f.latitude, s=sizes, color="#087f8c" if name != "Sampling candidates" else "#B37A21",
                   edgecolors="white", linewidth=.4, alpha=.75, label=label, zorder=4)
    ax.set(title="Water study boundary and sampling locations", xlabel="Longitude (WGS84)", ylabel="Latitude (WGS84)")
    ax.set_aspect(1 / max(.1, math.cos(math.radians(run["study"]["lat"]))))
    ax.grid(alpha=.2); ax.ticklabel_format(style="plain", useOffset=False)
    from matplotlib.ticker import MaxNLocator
    ax.xaxis.set_major_locator(MaxNLocator(5)); ax.yaxis.set_major_locator(MaxNLocator(5))
    if ax.get_legend_handles_labels()[0]:
        ax.legend(fontsize=7, loc="upper center", bbox_to_anchor=(.5, -.17))
    ax.annotate("N", xy=(.95, .94), xytext=(.95, .82), xycoords="axes fraction", ha="center", arrowprops=dict(arrowstyle="->"))
    fig.text(.08, -.17, "Coordinate map; no basemap imagery. Points represent samples or unverified sampling candidates, not continuous concentrations.", fontsize=7)
    return figure_png(fig)


def report_scope_notes(run):
    study = run["study"]
    notes = [f"{study.get('waterbody_type', 'Water')} study: {study['area_km2']:.3f} km² around {study['lat']:.5f}, {study['lon']:.5f}. This local boundary is not automatically the entire waterbody or catchment.",
             "Satellite indices, field measurements and model forecasts have different spatial/temporal support and are reported separately."]
    from evidence import quality_findings
    notes += [row["message"] for row in quality_findings(run) if row["code"] != "local_area"]
    return notes
